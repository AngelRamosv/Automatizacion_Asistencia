"""Procesador de Premios de Asistencia (Bonos P15 y P16 - $115.4 c/u)."""
import os
import openpyxl
import pandas as pd


def extraer_bonos_concentrado(path_concentrado, col_name, expected_dates_dmy):
    """
    Genera el reporte de Percepciones Asistencia P15 y P16 desde el concentrado.
    Busca la columna col_name ('BONO' o 'BONO AUTORIZADO').
    Todos los agentes con OK en esa columna reciben P15 y P16 con monto fijo 115.4.
    """
    fecha_fin_semana = expected_dates_dmy[6] if len(expected_dates_dmy) > 6 else expected_dates_dmy[-1]
    
    cols = [
        "Clave trabajador", "Per/Ded", "Num Per/Ded", "Num credito", "Aplica Destaj", "Aplicacion",
        "Fecha inicio", "Fecha fin", "Monto o formula"
    ]
    
    if not os.path.exists(path_concentrado):
        return pd.DataFrame(columns=cols)

    wb = openpyxl.load_workbook(path_concentrado, data_only=True)
    ws = wb.active

    bono_col = None
    for r_h in [1, 2]:
        for c in range(1, ws.max_column + 1):
            val = str(ws.cell(row=r_h, column=c).value or "").strip().upper()
            if val == col_name.upper():
                bono_col = c
                break
        if bono_col:
            break

    if bono_col is None:
        wb.close()
        return pd.DataFrame(columns=cols)

    movimientos = []
    for r in range(2, ws.max_row + 1):
        noi_val = ws.cell(row=r, column=3).value
        if noi_val is None:
            continue
        try:
            noi_code = int(float(str(noi_val).strip()))
        except (ValueError, TypeError):
            continue
        if noi_code <= 0:
            continue

        bono_val = str(ws.cell(row=r, column=bono_col).value or "").strip().upper()
        if bono_val == "OK":
            movimientos.append({
                "Clave trabajador": str(noi_code), "Per/Ded": "P", "Num Per/Ded": 15,
                "Num credito": 0, "Aplica Destaj": "N", "Aplicacion": 4,
                "Fecha inicio": fecha_fin_semana, "Fecha fin": fecha_fin_semana, "Monto o formula": 115.4
            })
            movimientos.append({
                "Clave trabajador": str(noi_code), "Per/Ded": "P", "Num Per/Ded": 16,
                "Num credito": 0, "Aplica Destaj": "N", "Aplicacion": 4,
                "Fecha inicio": fecha_fin_semana, "Fecha fin": fecha_fin_semana, "Monto o formula": 115.4
            })

    wb.close()
    
    df_premios = pd.DataFrame(movimientos)
    if not df_premios.empty:
        df_premios = df_premios[cols].sort_values("Clave trabajador").reset_index(drop=True)
    else:
        df_premios = pd.DataFrame(columns=cols)

    return df_premios
