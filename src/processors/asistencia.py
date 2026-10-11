"""Procesador de asistencia: IZZI VTAS y NOMINA POR SUPERVISOR IZZI ATN."""
import datetime
from copy import copy
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from src.config import DAY_INDICES, DAY_NAMES, ATTENDANCE_CODES, CODE_HEADERS


def formatear_ventas(output_vtas, reported_monday):
    """
    Escribe y formatea fechas estaticas en IZZI VTAS CONCENTRADO.xlsx
    para evitar celdas vacias o desbordamientos '###'.
    """
    print("Escribiendo y formateando fechas estáticas en IZZI VTAS CONCENTRADO.xlsx...")
    wb_vtas = openpyxl.load_workbook(output_vtas)
    ws_vtas = wb_vtas.active
    date_cols = ['I', 'K', 'M', 'O', 'Q', 'S', 'U']
    for idx, col in enumerate(date_cols):
        col_date = reported_monday + datetime.timedelta(days=idx)
        cell = ws_vtas[f"{col}1"]
        cell.value = col_date
        cell.number_format = 'dd/mm/yyyy'
        ws_vtas.column_dimensions[col].width = 12.0
    wb_vtas.save(output_vtas)
    wb_vtas.close()
    print("IZZI VTAS CONCENTRADO.xlsx creado con éxito con estilos y colores originales.")


def procesar_nomina_principal(output_principal, expected_dates_dmy):
    """
    Valida y formatea NOMINA POR SUPERVISOR IZZI ATN.xlsx:
    - Valida y corrige fechas en fila 1.
    - Extrae el listado de agentes.
    - Agrega columnas de conteo con formulas COUNTIF y estilos.
    - Agrega fila final de TOTAL GENERAL con formula SUM.
    """
    print("Cargando NOMINA POR SUPERVISOR IZZI ATN.xlsx para procesamiento de asistencia...")
    wb_principal = openpyxl.load_workbook(output_principal)
    ws = wb_principal["NOMINA POR SUPERVISOR IZZI ATN"]
    ws.views.sheetView[0].showGridLines = True
    
    max_row = ws.max_row
    print(f"Filas detectadas en la hoja original: {max_row}")
    
    corrections_made = []
    for idx, col_idx in enumerate(DAY_INDICES):
        cell = ws.cell(row=1, column=col_idx + 1)
        old_val = str(cell.value or "").strip()
        new_val = expected_dates_dmy[idx]
        if old_val != new_val:
            warning_msg = f"{DAY_NAMES[idx]} (Columna {col_idx+1}): Se leyó '{old_val}', se corrigió a '{new_val}'"
            print(f"!ADVERTENCIA! {warning_msg}")
            corrections_made.append(warning_msg)
            cell.value = new_val

    # Parsear registros de agentes
    agents_records = []
    for r in range(3, max_row + 1):
        noi_val = ws.cell(row=r, column=4).value  # Columna D es NOI
        if noi_val is None or str(noi_val).strip() == "" or str(noi_val).strip().lower() == "noi":
            continue
            
        try:
            noi_code = int(float(str(noi_val).strip()))
        except ValueError:
            continue
            
        supervisor = str(ws.cell(row=r, column=2).value or "").strip()  # Columna B
        puesto = str(ws.cell(row=r, column=5).value or "").strip()      # Columna E
        jornada = str(ws.cell(row=r, column=6).value or "").strip()     # Columna F
        nombre = str(ws.cell(row=r, column=7).value or "").strip()      # Columna G
        apellido = str(ws.cell(row=r, column=8).value or "").strip()    # Columna H
        
        day_codes = []
        for col_idx in DAY_INDICES:
            code = str(ws.cell(row=r, column=col_idx + 1).value or "").strip().upper()
            day_codes.append(code if code else "")
            
        agents_records.append({
            'row_index_1based': r,
            'noi_code': noi_code,
            'supervisor': supervisor,
            'puesto': puesto,
            'jornada': jornada,
            'nombre': nombre,
            'apellido': apellido,
            'day_codes': day_codes
        })
        
    print(f"Total de agentes identificados para procesamiento: {len(agents_records)}")
    
    # Agregar columnas de conteo y formulas COUNTIF
    ref_header_cell = ws.cell(row=2, column=5)  # 'EMPLEADO POR NOMBRE'
    ref_font = ref_header_cell.font
    ref_fill = ref_header_cell.fill
    ref_border = ref_header_cell.border
    
    start_col = 22  # Columna V
    
    # Encabezados en Fila 2
    for idx, header_text in enumerate(CODE_HEADERS):
        col = start_col + idx
        cell = ws.cell(row=2, column=col, value=header_text)
        if ref_font:
            cell.font = copy(ref_font)
        if ref_fill:
            cell.fill = copy(ref_fill)
        if ref_border:
            cell.border = copy(ref_border)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        
    # Fondo/borde de Fila 1 para nuevas columnas
    for idx in range(len(CODE_HEADERS)):
        col = start_col + idx
        cell_r1 = ws.cell(row=1, column=col, value="")
        if ref_fill:
            cell_r1.fill = copy(ref_fill)
        if ref_border:
            cell_r1.border = copy(ref_border)
            
    thin_border = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )
    
    # Formulas COUNTIF por agente
    for agent in agents_records:
        r = agent['row_index_1based']
        for idx, code in enumerate(ATTENDANCE_CODES):
            col = start_col + idx
            col_letter_start = "H"
            col_letter_end = "T"
            
            if code == 'A':
                formula = f'=COUNTIF({col_letter_start}{r}:{col_letter_end}{r}, "A") + COUNTIF({col_letter_start}{r}:{col_letter_end}{r}, "R")'
            else:
                formula = f'=COUNTIF({col_letter_start}{r}:{col_letter_end}{r}, "{code}")'
                
            cell = ws.cell(row=r, column=col, value=formula)
            cell.font = Font(name="Calibri", size=11)
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
            
            # Resaltado condicional sutil
            code_count = agent['day_codes'].count(code) + (agent['day_codes'].count('R') if code == 'A' else 0)
            if code_count > 0:
                if code in ['F', 'SUSP']:
                    cell.fill = PatternFill(start_color="FCE4D6", fill_type="solid")
                    cell.font = Font(name="Calibri", size=11, bold=True)
                elif code == 'A':
                    cell.fill = PatternFill(start_color="E2EFDA", fill_type="solid")
                elif code == 'V':
                    cell.fill = PatternFill(start_color="FFF2CC", fill_type="solid")
                elif code == 'INC':
                    cell.fill = PatternFill(start_color="E8DAEF", fill_type="solid")
                    
    # Fila de TOTAL GENERAL
    total_row = max_row + 1
    ws.cell(row=total_row, column=6, value="TOTAL GENERAL").font = Font(name="Calibri", size=11, bold=True)
    ws.cell(row=total_row, column=6).alignment = Alignment(horizontal="right", vertical="center")
    
    for idx, code in enumerate(ATTENDANCE_CODES):
        col = start_col + idx
        col_letter = get_column_letter(col)
        cell = ws.cell(row=total_row, column=col, value=f"=SUM({col_letter}3:{col_letter}{total_row-1})")
        cell.font = Font(name="Calibri", size=11, bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = Border(top=Side(style='thin', color='000000'), bottom=Side(style='double', color='000000'))
        cell.fill = PatternFill(start_color="D9E1F2", fill_type="solid")
        
    # Ajuste de ancho de columnas
    for col_idx in range(start_col, start_col + len(ATTENDANCE_CODES)):
        col_letter = get_column_letter(col_idx)
        max_len = 0
        for r in range(2, total_row + 1):
            val = ws.cell(row=r, column=col_idx).value
            if val is not None:
                max_len = max(max_len, len(str(val)))
        ws.column_dimensions[col_letter].width = max(max_len + 3, 14)
        
    ws.row_dimensions[2].height = 28
    ws.row_dimensions[total_row].height = 22
    
    wb_principal.save(output_principal)
    wb_principal.close()
    print("NOMINA POR SUPERVISOR IZZI ATN.xlsx actualizado con éxito.")
    
    return agents_records, corrections_made
