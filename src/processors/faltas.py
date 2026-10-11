"""Procesador de extraccion y formato de incidencias y faltas."""
import os
import openpyxl
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
import pandas as pd

from src.config import ABSENCE_CODES, DATE_COLS_CONCENTRADO, RETARDOS_COL


def extraer_faltas_concentrado(path, expected_dates_dmy):
    """
    Lee el concentrado directamente fila por fila.
    NOI en columna C (3). Las fechas y codigos estan en columnas impares (17, 19, 21, 23, 25, 27, 29).
    Genera UNA fila por cada dia con incidencia valida (Tipo falta = 2).
    Aplica la regla de 3 retardos = 1 falta adicional.
    """
    rows = []
    if not os.path.exists(path):
        print(f"  ADVERTENCIA: No se encontró {path}")
        return rows
    
    wb_c = openpyxl.load_workbook(path, data_only=True)
    ws_c = wb_c.active
    
    for r in range(2, ws_c.max_row + 1):
        noi_val = ws_c.cell(row=r, column=3).value  # NOI en col C
        if noi_val is None:
            continue
        try:
            noi_code = int(float(str(noi_val).strip()))
        except (ValueError, TypeError):
            continue
        if noi_code <= 0:
            continue
        
        for i, code_col in enumerate(DATE_COLS_CONCENTRADO):
            date_str = expected_dates_dmy[i] if i < len(expected_dates_dmy) else ""
            if not date_str:
                continue
            
            code_raw = str(ws_c.cell(row=r, column=code_col).value or "").strip().upper()
            
            # Ignorar descansos, asistencias y blancos
            if code_raw in ('D', '', 'NONE', 'A', 'R', 'V'):
                continue
            if code_raw not in ABSENCE_CODES:
                continue
            
            rows.append({
                'Clave trab': noi_code,
                'Tipo falta': 2,
                '% falta': 100,
                'Certificado': 0,
                '% pagado': 0,
                'Fecha inic': date_str,
                'Fecha fin': date_str
            })
            
        # --- REGLA RETARDOS: si el agente tiene exactamente 3 retardos, se agrega 1 falta extra ---
        retardos_val = ws_c.cell(row=r, column=RETARDOS_COL).value
        try:
            retardos = int(float(str(retardos_val).strip())) if retardos_val is not None else 0
        except (ValueError, TypeError):
            retardos = 0
        
        if retardos == 3:
            fechas_con_falta = {row['Fecha inic'] for row in rows if row['Clave trab'] == noi_code}
            fecha_extra = None
            for fecha_libre in expected_dates_dmy:
                if fecha_libre not in fechas_con_falta:
                    fecha_extra = fecha_libre
                    break
            if fecha_extra:
                rows.append({
                    'Clave trab': noi_code,
                    'Tipo falta': 2,
                    '% falta': 100,
                    'Certificado': 0,
                    '% pagado': 0,
                    'Fecha inic': fecha_extra,
                    'Fecha fin': fecha_extra
                })
    
    wb_c.close()
    return rows


def guardar_y_estilizar_faltas(df_to_save, out_path):
    """Guarda el DataFrame de faltas en formato Excel con estilos de celdas y anchos de columna."""
    df_to_save.to_excel(out_path, index=False, sheet_name="Faltas")
    wb_f = openpyxl.load_workbook(out_path)
    ws_f = wb_f["Faltas"]
    ws_f.views.sheetView[0].showGridLines = True
    for col in ws_f.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row > 1 and col_letter in ['F', 'G'] and cell.value is not None:
                cell.alignment = Alignment(horizontal="center")
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws_f.column_dimensions[col_letter].width = max(max_len + 3, 14)
    wb_f.save(out_path)
    wb_f.close()
