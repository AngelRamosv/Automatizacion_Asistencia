"""Procesador de calculo de Prima Dominical (P19 - 25% salario diario)."""
import os
import openpyxl
import pandas as pd
from src.config import SUNDAY_CODE_COL


def generar_prima_dominical(df_empleados, expected_dates_dmy, output_path_cyber, output_path_wincall, path_cyber_conc, path_wincall_conc):
    """
    Genera DOS reportes de Percepciones (Prima Dominical P19), uno por empresa.
    Se guia EXCLUSIVAMENTE de los concentrados de ambas empresas.
    Solo incluye agentes con A o R en el dia domingo (columna 29 / AC).
    Calcula: round(sueldo_diario * 0.25, 2)
    Inyecta la hoja 'Prima_Dominical' en los archivos de incidencias.
    """
    fecha_domingo = expected_dates_dmy[6] if len(expected_dates_dmy) > 6 else ""
    cols = [
        'Clave trabajador', 'Per/Ded', 'Num Per/Ded', 'Num credito', 'Aplica Destaj', 'Aplicacion',
        'Fecha inicio', 'Fecha fin', 'Monto o formula'
    ]
    
    def extraer_nois_con_domingo(path):
        nois_domingo = set()
        if not os.path.exists(path):
            print(f"  ADVERTENCIA Percepciones: No se encontro {path}")
            return nois_domingo
        
        wb_c = openpyxl.load_workbook(path, data_only=True)
        ws_c = wb_c.active
        
        for r in range(2, ws_c.max_row + 1):
            noi_val = ws_c.cell(row=r, column=3).value
            if noi_val is None:
                continue
            try:
                noi_code = int(float(str(noi_val).strip()))
            except (ValueError, TypeError):
                continue
            if noi_code <= 0:
                continue
            code = str(ws_c.cell(row=r, column=SUNDAY_CODE_COL).value or "").strip().upper()
            if code in ('A', 'R'):
                nois_domingo.add(noi_code)
        
        wb_c.close()
        return nois_domingo
    
    def build_movimientos(nois_set):
        movimientos = []
        for noi_code in nois_set:
            clave = str(noi_code)
            empleado = df_empleados[df_empleados['Clave trabajador'].astype(str) == clave]
            if not empleado.empty:
                sueldo_diario = pd.to_numeric(empleado['Salario diario'].values[0], errors='coerce')
                if pd.notnull(sueldo_diario):
                    monto_prima = round(sueldo_diario * 0.25, 2)
                    movimientos.append({
                        'Clave trabajador': clave, 'Per/Ded': 'P', 'Num Per/Ded': 19,
                        'Num credito': 0, 'Aplica Destaj': 'N', 'Aplicacion': 4,
                        'Fecha inicio': fecha_domingo, 'Fecha fin': fecha_domingo,
                        'Monto o formula': monto_prima
                    })
        df = pd.DataFrame(movimientos)
        if not df.empty:
            return df[cols].sort_values('Clave trabajador').reset_index(drop=True)
        return pd.DataFrame(columns=cols)
    
    nois_cyber = extraer_nois_con_domingo(path_cyber_conc)
    nois_wincall = extraer_nois_con_domingo(path_wincall_conc)
    
    print(f"  -> {len(nois_cyber)} agentes CyberIdeas con asistencia el domingo")
    print(f"  -> {len(nois_wincall)} agentes WinCall con asistencia el domingo")
    
    df_cyber = build_movimientos(nois_cyber)
    df_wincall = build_movimientos(nois_wincall)
    
    with pd.ExcelWriter(output_path_cyber, engine='openpyxl', mode='a', if_sheet_exists='replace') as writer:
        df_cyber.to_excel(writer, index=False, sheet_name="Prima_Dominical")
    print(f"Percepciones_CyberIdeas.xlsx generado: {output_path_cyber}")
    
    with pd.ExcelWriter(output_path_wincall, engine='openpyxl', mode='a', if_sheet_exists='replace') as writer:
        df_wincall.to_excel(writer, index=False, sheet_name="Prima_Dominical")
    print(f"Percepciones_WinCall.xlsx generado: {output_path_wincall}")
